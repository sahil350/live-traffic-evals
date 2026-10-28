"""Build blind judge packets and answer keys for the WIDTH/DEPTH judges.

Usage: python3 -I build_judge_packets.py [--prompt data_collection/<file> --tag v3] E9 E20 ...   (case ids)
  Default prompt is judge_system_prompt_v2.txt and packets go to packets/. With --tag T they go to packets_T/.

Reads:
  outputs/openresearcher_reference_checklists_batch1_equivalence_v1.1_core_extra.xlsx
      Cases, WIDTH Equivalence v1, Core-Extra Rubric
  outputs/openresearcher_checklist_atomicity_revision_draft_v1.xlsx   (satisfying paths for split items)
  outputs/openresearcher_run/batch_results_free.jsonl                  (the answers and trajectories)
  data_collection/judge_system_prompt_v2.txt                           (the judge prompt)

Writes:
  outputs/judge_runs/packets/<case>.md   what a judge sees (tiers, weights and reference grades are hidden)
  outputs/judge_runs/key/<case>.json     tier and weight per item, used only for scoring afterwards
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
PROD = OUT / "openresearcher_reference_checklists_batch1_equivalence_v1.1_core_extra.xlsx"
DRAFT = OUT / "openresearcher_checklist_atomicity_revision_draft_v1.xlsx"
RESULTS = OUT / "openresearcher_run" / "batch_results_free.jsonl"
PROMPT = ROOT / "data_collection" / "judge_system_prompt_v2.txt"

MARK = re.compile(r"[【\[](\d+)†L(\d+)(?:-L(\d+))?[】\]]")


def sheet(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:]]


def excerpts(answer, messages):
    """Cited lines (plus two lines of context) for every line-range marker in the answer."""
    tool = {}
    for m in messages:
        if m.get("role") == "tool":
            h = re.match(r"\[(\d+)\] ", m.get("content") or "")
            if h and h.group(1) not in tool:
                tool[h.group(1)] = m["content"]
    ranges = {}
    for c, s, e in MARK.findall(answer):
        s = int(s)
        e = int(e) if e else s
        ranges.setdefault(c, []).append((max(0, s - 2), e + 2))
    blocks = []
    for c in sorted(ranges, key=int):
        if c not in tool:
            blocks.append(f"[{c}] SOURCE NOT FOUND IN THE RETRIEVED MATERIAL")
            continue
        text = tool[c]
        lines = {int(m.group(1)): m.group(2) for m in re.finditer(r"^L(\d+): ?(.*)$", text, re.M)}
        merged = []
        for s, e in sorted(ranges[c]):
            if merged and s <= merged[-1][1] + 1:
                merged[-1][1] = max(merged[-1][1], e)
            else:
                merged.append([s, e])
        body = []
        for s, e in merged:
            for n in range(s, e + 1):
                if n in lines:
                    body.append(f"L{n}: {lines[n]}")
            body.append("...")
        blocks.append(text.split("\n")[0][:200] + "\n" + "\n".join(body))
    return "\n\n".join(blocks) if blocks else "(the answer has no line-range citations)"


def main(case_ids, prompt=PROMPT, tag="", results_path=RESULTS):
    prod = openpyxl.load_workbook(PROD, read_only=True)
    draft = openpyxl.load_workbook(DRAFT, read_only=True)
    cases = {str(r["case_id"]): r for r in sheet(prod, "Cases")}
    equiv = {(str(r["case_id"]), r["item_id"]): r for r in sheet(prod, "WIDTH Equivalence v1")}
    split_equiv = {(str(r["case_id"]), r["item_id"]): r for r in sheet(draft, "Proposed_Splits_WIDTH")}
    rubric = [r for r in sheet(prod, "Core-Extra Rubric")]
    results = {str(json.loads(l)["qid"]): json.loads(l) for l in open(results_path)}
    system = Path(prompt).read_text()
    pdir = OUT / "judge_runs" / (f"packets_{tag}" if tag else "packets")
    pdir.mkdir(parents=True, exist_ok=True)
    (OUT / "judge_runs" / "key").mkdir(parents=True, exist_ok=True)

    for q in case_ids:
        c = cases[q]
        res = results[q]
        answer = [m for m in res["messages"] if m["role"] == "assistant"][-1]["content"]
        items = [r for r in rubric if str(r["case_id"]) == q and r["tier"] != "na"]
        parts = [
            "# JUDGE INSTRUCTIONS\n\n" + system.strip(),
            f"\n\n# CASE {q}\n\nEvaluation date: {c.get('as_of_date')}\nEvaluation assumption: {c.get('evaluation_assumption')}\n",
            "\n## QUESTION\n\n" + str(c["prompt"]).strip(),
            "\n\n## ANSWER BEING GRADED\n\n" + answer.strip(),
            "\n\n## CITED SOURCE EXCERPTS\n\n" + excerpts(answer, res["messages"]),
            "\n\n## CHECKLIST ITEMS (grade every item listed)\n",
        ]
        key = {"case_id": q, "items": []}
        for r in items:
            iid = r["item_id"]
            e = split_equiv.get((q, iid)) if r.get("parent_item_id") else equiv.get((q, iid))
            req = r.get("core_requirement") or r["requirement"]
            block = [f"\n### {iid} [{r['item_type']}]", f"Requirement: {req}"]
            if e and e.get("satisfying_paths"):
                block.append("Satisfying paths:\n" + str(e["satisfying_paths"]))
                block.append("Conditions (numbered to match the paths):\n" + str(e.get("conditions_aligned_by_number")))
                block.append(f"open_set: {e.get('open_set')}")
                block.append(f"accept_other_if: {e.get('accept_other_if')}")
            parts.append("\n".join(block))
            key["items"].append({"item_id": iid, "tier": r["tier"], "tier_call": r["tier_call"], "weight": int(r["weight"]),
                                 "item_type": r["item_type"], "used_core_requirement": bool(r.get("core_requirement"))})
        (pdir / f"{q}.md").write_text("\n".join(parts) + "\n")
        (OUT / "judge_runs" / "key" / f"{q}.json").write_text(json.dumps(key, indent=1))
        print(q, "items to judge:", len(items), "tiers:", sorted({i["tier"] for i in key["items"]}))


if __name__ == "__main__":
    args = sys.argv[1:]
    prompt, tag, results_path = PROMPT, "", RESULTS
    if "--prompt" in args:
        i = args.index("--prompt"); prompt = ROOT / args[i + 1]; del args[i:i + 2]
    if "--tag" in args:
        i = args.index("--tag"); tag = args[i + 1]; del args[i:i + 2]
    if "--results" in args:
        i = args.index("--results"); results_path = OUT / "openresearcher_run" / args[i + 1]; del args[i:i + 2]
    main(args, prompt, tag, results_path)
